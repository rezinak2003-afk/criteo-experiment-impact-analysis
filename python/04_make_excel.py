"""Builds the Excel artifact. Create the PivotTable yourself in Excel (openpyxl cannot create pivots)."""
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from config import ROOT
arm = pd.read_csv(ROOT / "outputs" / "arm_summary.csv")
res = pd.read_csv(ROOT / "outputs" / "uplift_results.csv")
wb = Workbook(); ws = wb.active; ws.title = "Arm_Summary"
ws.append(list(arm.columns))
for r in arm.itertuples(index=False): ws.append(list(r))
rc = wb.create_sheet("Reconciliation")
rc.append(["arm", "SQL visit_rate", "Excel visit_rate (n_visits/n_users)", "XLOOKUP visit_rate", "INDEX/MATCH fallback", "difference"])
for i, a in enumerate(arm["arm"], start=2):
    rc.append([a, float(arm.loc[i - 2, "visit_rate"]), f"=INDEX(Arm_Summary!E:E,MATCH(A{i},Arm_Summary!A:A,0))/INDEX(Arm_Summary!C:C,MATCH(A{i},Arm_Summary!A:A,0))",
               f'=IFERROR(_xlfn.XLOOKUP(A{i},Arm_Summary!A:A,Arm_Summary!F:F),"missing")',
               f'=IFERROR(INDEX(Arm_Summary!F:F,MATCH(A{i},Arm_Summary!A:A,0)),"missing")', f"=ABS(B{i}-C{i})"])
k = wb.create_sheet("KPI_Summary"); k.append(list(res.columns))
for r in res.itertuples(index=False): k.append(list(r))
k.append([]); k.append(["Note: values come from python/03_analysis.py; insert a PivotTable on Arm_Summary yourself."])
for sh in wb: 
    for c in sh[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F3A5F")
wb.save(ROOT / "excel" / "reconciliation_and_summary.xlsx"); print("saved excel/reconciliation_and_summary.xlsx")
